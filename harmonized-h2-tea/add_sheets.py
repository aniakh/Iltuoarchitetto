#!/usr/bin/env python3
"""Second pass: incremental-NPV view, Validation_Tests, Corrections_Log, README banner.
Also forces full recalculation on open so no stale cached value is shipped."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BLUE = Font(name='Arial', size=10, color='FF0000FF')
BLACK = Font(name='Arial', size=10, color='FF000000')
GREEN = Font(name='Arial', size=10, color='FF008000')
HDR = Font(name='Arial', size=11, bold=True, color='FF1F6F5C')
SUB = Font(name='Arial', size=10, bold=True, color='FF1F6F5C')
REDB = Font(name='Arial', size=10, bold=True, color='FFC00000')
WHITEB = Font(name='Arial', size=11, bold=True, color='FFFFFFFF')
HDRFILL = PatternFill('solid', fgColor='FFE8F0EC')
WARNFILL = PatternFill('solid', fgColor='FFFFF3CD')
BANNER = PatternFill('solid', fgColor='FFC00000')
wrap = Alignment(wrap_text=True, vertical='top')

FILES = {
    'UCI_CentralPlant_Harmonized_TEA_Model.xlsx': 16,
    'UCI_FineArts_Harmonized_TEA_Model.xlsx': 14,
    'UCI_Microturbine_Harmonized_TEA_Model.xlsx': 15,
    'PGE_Harmonized_TEA_Model.xlsx': 13,
}

CORRECTIONS = [
    ("1. Units (LHV)", "H2 energy content now LHV-consistent (=33.3 kWh/kg x 0.0036 = 0.11988 GJ/kg). "
        "The old 0.113745 was ~5.4% low and overstated H2 mass demand, cost, storage, compression and credits."),
    ("2. IRA 45V", "Removed from the base case. Behind an eligibility toggle; earned by the qualified H2 producer, "
        "capped at 10 years in the cashflow (was applied ~50 yrs to the end user)."),
    ("3. LCFS", "Removed from the base case. Behind an eligibility toggle; eligible price repriced to CARB-recent "
        "~$66/tCO2e (was $150). Stationary CHP/kiln/transport eligibility must be documented before use."),
    ("4. Offtake revenue", "The $4/kg H2 'production revenue' is OFF in the base case for end uses "
        "(a consumer cannot pay for H2 and also book production revenue). Use incremental NPV / avoided-cost framing."),
    ("5. CCS", "Now an explicit on/off option (default OFF). Previously CCS cost was charged in EVERY scenario, "
        "including 0% H2, inflating LCOH."),
    ("6. Efficiency sign", "Combustion efficiency rises slightly with H2 everywhere (Thermoflex); the Cashflow sheet "
        "previously used the opposite sign (a derate) vs the Fuel_Cost sheet."),
    ("7. Availability", "LCOH now uses the SAME derived 50-yr mean availability as the cashflow (was a fixed 0.97 in LCOH "
        "vs ~0.955 derived in the cashflow)."),
    ("8. LCOH at 0%", "Shown as N/A (no hydrogen exists at 0% blend); removes the -$1.35/kg artefact at 0%."),
    ("9. Network units", "TEA_Network total no longer adds $/kg layers to a $/MWh integration term; the stale cached "
        "total is refreshed on recalculation."),
    ("10. Incremental NPV", "NPV_Comparison now reports NPV incremental to the 0% fossil baseline; the 0% incremental "
        "NPV is zero by construction. Absolute 'project' NPV is retained but flagged non-decision-grade."),
]

PRELIMINARY = [
    ("Hydraulics / compressors", "N_comp = MAX(1, distance/spacing x flow-factor) is a screening PROXY, not compressible "
        "pipeline hydraulics. It does not verify delivery pressure or compute compressor duty from pressure ratio. "
        "Do not treat compressor counts or the 3.1 kWh/kg adder as design values."),
    ("Reliability PoF", "Component PoF are cumulative 'at least one functional failure' probabilities, NOT catastrophic "
        "integrity risk, and must not be read alongside rupture probability as equivalent consequence."),
    ("Lifecycle degradation", "The lifecycle uses a CONSTANT hazard (exponential), not an age-dependent Weibull/renewal "
        "model, despite 'PoF rises with age' language. Residual-life claims are indicative only."),
    ("Pipe integrity", "The FCGR step multipliers are a screening failure-rate multiplier model, NOT probabilistic "
        "fracture mechanics. 'Integrity saturates at 20%' is an artefact of the imposed step function, not a physical result."),
    ("CAPEX scaling", "Pipeline $/mile, the $2M skid and compressor costs are not size/diameter/pressure/location-specific. "
        "Small low-utilisation loads (e.g. Fine Arts) show LCOH dominated by fixed-cost dilution, not technology cost."),
    ("Equipment limits", "The 6.89-bar H2 partial-pressure rule is an API-617 centrifugal-compressor concern, not a "
        "universal component acceptance criterion. Replacement CAPEX should follow component-specific qualification."),
    ("Prices", "H2 $2/kg is optimistic for clean delivered H2 (DOE ~$5-7/kg plant-gate); NG $4.74/GJ and electricity "
        "$0.235/kWh should be validated against actual bills/tariffs. Prices are screening placeholders."),
    ("Reconciliation", "A separate SoCalGas workbook referenced by the review was not provided; only 4 workbooks + the "
        "dashboard are in this set. Report/dashboard/workbook reconciliation to a single engine remains outstanding."),
]

def add_banner_border(ws, r1, r2, c1, c2):
    thin = Side(style='thin', color='FFBBBBBB')
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = Border(thin, thin, thin, thin)

def process(fn, avail_row):
    wb = openpyxl.load_workbook(fn)

    # ---- Incremental NPV column on NPV_Comparison block A --------------------
    npv = wb['NPV_Comparison']
    # Fix efficiency sign inline in any C-column closed form
    for row in npv.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and '1-Global_Assumptions!$B$64' in cell.value:
                cell.value = cell.value.replace('1-Global_Assumptions!$B$64', '1+Global_Assumptions!$B$64')
    npv['D6'] = 'Absolute NPV 50yr ($) — non-decision-grade'
    npv['D6'].font = REDB
    npv['F6'] = 'Incremental NPV vs 0% ($) — decision-relevant'
    npv['F6'].font = SUB
    for r in range(7, 15):  # 0,5,10,20,30,50,75,100 near-retrofit block
        npv.cell(r, 6, '=D%d-$D$7' % r).font = BLACK
    npv['A2'] = ('Screening NPV. Lead with the INCREMENTAL column (vs the 0% fossil baseline); the 0% incremental NPV '
                 'is zero by construction. Absolute NPV is retained for continuity only and is non-decision-grade '
                 '(end-use benefit = avoided electricity/heat, not modelled here; transmission uses a merchant revenue proxy).')
    npv['A2'].font = BLACK; npv['A2'].alignment = wrap

    # ---- Validation_Tests sheet --------------------------------------------
    if 'Validation_Tests' in wb.sheetnames:
        del wb['Validation_Tests']
    vt = wb.create_sheet('Validation_Tests')
    vt.column_dimensions['A'].width = 4
    vt.column_dimensions['B'].width = 34
    vt.column_dimensions['C'].width = 16
    vt.column_dimensions['D'].width = 66
    vt['A1'] = 'MINIMUM VALIDATION TESTS (auto-checked; recalculates on open)'
    vt['A1'].font = HDR
    vt['A3'] = '#'; vt['B3'] = 'Test'; vt['C3'] = 'Result'; vt['D3'] = 'Basis / note'
    for c in 'ABCD':
        vt['%s3' % c].font = SUB; vt['%s3' % c].fill = HDRFILL
    tests = [
        ('Zero-blend', '=IF(AND(Blend_Properties!$F$6=0,Scenario_Results!$G$4=0,Scenario_Results!$L$4=0),"PASS","FAIL")',
         'At 0% H2: demand, H2 CAPEX and CCS cost all zero.'),
        ('Baseline (incremental NPV=0)', '=IF(ABS(NPV_Comparison!$F$7)<1,"PASS","FAIL")',
         'Incremental NPV at 0% blend is zero by construction.'),
        ('Units (LHV consistent)', '=IF(ABS(Global_Assumptions!$B$25-Global_Assumptions!$B$24*0.0036)<0.000001,"PASS","FAIL")',
         'H2 GJ/kg equals LHV_kWh x 0.0036.'),
        ('Units (no $/kg + $/MWh sum)', '=IF(ABS(TEA_Network!$B$8-(TEA_Network!$B$4+TEA_Network!$B$5+TEA_Network!$B$6))<0.000001,"PASS","FAIL")',
         'Network total excludes the $/MWh integration term.'),
        ('Policy: 45V capped at 10 yr', '=IF(Global_Assumptions!$B$86=10,"PASS","FAIL")',
         '45V duration input = 10 years; cashflow gates 45V to years 1..10.'),
        ('Credit: LCFS=0 in base case', '=IF(Global_Assumptions!$B$20=0,"PASS","FAIL")',
         'LCFS is zero unless eligibility toggle = 1.'),
        ('Hydraulic: outlet pressure OK', '="PRELIMINARY"',
         'Not modelled — N_comp is a screening proxy; compressible hydraulics is future work.'),
        ('Availability consistency', '=IF(ABS(Global_Assumptions!$B$11-Lifecycle_50yr!$B$%d)<0.000001,"PASS","FAIL")' % avail_row,
         'LCOH and cashflow use the same derived 50-yr mean availability.'),
        ('Reconciliation (dashboard/report)', '="MANUAL"',
         'Reconcile to a single engine; SoCalGas workbook not provided in this set.'),
        ('Manual benchmark (0/20/100%)', '="MANUAL"',
         'Independently reproduce the 0%, 20% and 100% cases before quoting results.'),
    ]
    r = 4
    for i, (name, formula, basis) in enumerate(tests, 1):
        vt.cell(r, 1, i).font = BLACK
        vt.cell(r, 2, name).font = BLACK
        cc = vt.cell(r, 3, formula); cc.font = BLACK
        vt.cell(r, 4, basis).font = BLACK; vt.cell(r, 4).alignment = wrap
        r += 1
    vt.cell(r + 1, 2, 'PASS/FAIL cells recalculate when the file is opened in Excel/LibreOffice.').font = BLACK

    # ---- Corrections_Log sheet ---------------------------------------------
    if 'Corrections_Log' in wb.sheetnames:
        del wb['Corrections_Log']
    cl = wb.create_sheet('Corrections_Log')
    cl.column_dimensions['A'].width = 24
    cl.column_dimensions['B'].width = 96
    cl['A1'] = 'CORRECTIONS APPLIED (technical review response)'
    cl['A1'].font = HDR
    cl['A2'] = ('PRELIMINARY / SCREENING ONLY — NON-DECISION-GRADE. These corrections fix boundary, unit, policy and '
                'consistency errors so the base case is defensible; they do not turn the suite into an investment-grade model.')
    cl['A2'].font = REDB; cl['A2'].alignment = wrap; cl.merge_cells('A2:B2')
    r = 4
    cl.cell(r, 1, 'Corrected in this revision').font = SUB
    cl.cell(r, 1).fill = HDRFILL; cl.cell(r, 2).fill = HDRFILL
    r += 1
    for name, desc in CORRECTIONS:
        cl.cell(r, 1, name).font = BLACK
        cl.cell(r, 2, desc).font = BLACK; cl.cell(r, 2).alignment = wrap
        r += 1
    r += 1
    cl.cell(r, 1, 'Still preliminary / not rebuilt').font = SUB
    cl.cell(r, 1).fill = WARNFILL; cl.cell(r, 2).fill = WARNFILL
    r += 1
    for name, desc in PRELIMINARY:
        cl.cell(r, 1, name).font = REDB
        cl.cell(r, 2, desc).font = BLACK; cl.cell(r, 2).alignment = wrap
        r += 1

    # ---- README preliminary banner -----------------------------------------
    rd = wb['README']
    rd.insert_rows(1, 4)
    rd['A1'] = 'PRELIMINARY — SCREENING ONLY, NON-DECISION-GRADE (revised per technical review)'
    rd['A1'].font = WHITEB; rd['A1'].fill = BANNER
    rd['A2'] = ('Base case now excludes LCFS, IRA 45V, CCS and $/kg offtake revenue (all behind explicit toggles in '
                'Global_Assumptions); H2 energy basis is LHV-consistent; NPV is reported incremental to the 0% baseline. '
                'See the Corrections_Log and Validation_Tests sheets. Do NOT use LCOH/NPV/abatement figures for investment '
                'decisions or cross-site ranking.')
    rd['A2'].font = BLACK; rd['A2'].alignment = wrap
    rd.merge_cells('A1:B1'); rd.merge_cells('A2:B2')

    # ---- Force full recalc on open (refresh all cached values) --------------
    try:
        wb.calc_properties.fullCalcOnLoad = True
    except Exception:
        pass

    wb.save(fn)
    print('added sheets ->', fn)

if __name__ == '__main__':
    for fn, ar in FILES.items():
        process(fn, ar)
